"""実機テスト用の pptx を作る: python3 test/fixtures/make_test_decks.py [出力先フォルダー]

どのファイルもノートに ```mermaid ブロックを書いただけの状態で、図はまだ入っていない。
PowerPoint で開いて「全スライドを更新」を押すところからテストする（テスト計画 D・E）。

  perf-100x1.pptx    100 枚、各 1 図（P1〜P4）
  perf-300x1.pptx    300 枚、各 1 図（P8）
  perf-1x20.pptx     1 枚に 20 図（P5）
  perf-big.pptx      ノード 200 個のフローチャートと、メッセージ 100 個のシーケンス図（P6）
  perf-heavy.pptx    100 枚、各 1 図＋写真相当の画像で約 50MB（P7）
  compat-samples.pptx  パネルのサンプル集を 1 枚に 1 つずつ（E: PowerPoint 以外での表示）

python-pptx が必要（pip install python-pptx）。
"""
import os
import re
import struct
import sys
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]

# P1 は「図なし」のスライドで測るため、ノートは図の種類を回すだけにして、図の有無は枚数で調整する
ROTATION = ["flowchart", "sequence", "class", "state"]


def load_samples():
    """src/core/samples.ts から id とコードを読む（パネルのサンプル集と同じものを使う）"""
    text = (ROOT / "src/core/samples.ts").read_text(encoding="utf-8")
    pairs = re.findall(r'id: "([^"]+)",\s*label: "([^"]+)",\s*code: `([^`]*)`', text)
    if not pairs:
        raise SystemExit("samples.ts からサンプルを読めませんでした")
    return [{"id": i, "label": l, "code": c} for i, l, c in pairs]


def fence(block_id, code):
    return f"```mermaid id={block_id}\n{code}\n```"


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    return prs


def add_slide(prs, title, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def noise_png(path, width, height):
    """圧縮の効かない PNG（写真の代わり）。Pillow を使わずに書く"""
    raw = b"".join(b"\x00" + os.urandom(width * 3) for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw, 1)) + chunk(b"IEND", b""))


def big_flowchart(nodes=200):
    lines = ["flowchart TD"]
    for i in range(1, nodes):
        parent = (i - 1) // 3  # 3 分木
        lines.append(f"  N{parent}[ノード {parent}] --> N{i}[ノード {i}]")
    return "\n".join(lines)


def long_sequence(messages=100):
    lines = ["sequenceDiagram", "  participant U as ユーザー", "  participant A as アプリ", "  participant S as サーバー"]
    pairs = [("U", "A"), ("A", "S"), ("S", "A"), ("A", "U")]
    for i in range(messages):
        a, b = pairs[i % 4]
        lines.append(f"  {a}->>{b}: メッセージ {i + 1}")
    return "\n".join(lines)


def rotating_deck(samples, slides, per_slide=1):
    by_id = {s["id"]: s for s in samples}
    prs = new_deck()
    for n in range(slides):
        blocks = []
        for k in range(per_slide):
            sample = by_id[ROTATION[(n + k) % len(ROTATION)]]
            blocks.append(fence(f"d{k + 1}", sample["code"]))
        add_slide(prs, f"スライド {n + 1}", f"発表者メモ {n + 1}\n\n" + "\n\n".join(blocks))
    return prs


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "test-decks"
    out.mkdir(parents=True, exist_ok=True)
    samples = load_samples()

    rotating_deck(samples, 100).save(out / "perf-100x1.pptx")
    rotating_deck(samples, 300).save(out / "perf-300x1.pptx")
    rotating_deck(samples, 1, per_slide=20).save(out / "perf-1x20.pptx")

    prs = new_deck()
    add_slide(prs, "ノード 200 個のフローチャート", fence("big-flow", big_flowchart()))
    add_slide(prs, "メッセージ 100 個のシーケンス図", fence("long-seq", long_sequence()))
    prs.save(out / "perf-big.pptx")

    # 100 枚 × 約 0.5MB の画像で約 50MB。同じ画像は 1 つにまとめられるため、1 枚ずつ作る
    prs = rotating_deck(samples, 100)
    image = out / "_noise.png"
    for slide in prs.slides:
        noise_png(image, 420, 400)
        slide.shapes.add_picture(str(image), Inches(9), Inches(2), width=Inches(3.5))
    image.unlink()
    prs.save(out / "perf-heavy.pptx")

    prs = new_deck()
    for s in samples:
        add_slide(prs, s["label"], fence(s["id"], s["code"]))
    prs.save(out / "compat-samples.pptx")

    for f in sorted(out.glob("*.pptx")):
        print(f"{f.name:22} {f.stat().st_size / 1024 / 1024:6.1f} MB")


if __name__ == "__main__":
    main()
