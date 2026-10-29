"""テスト用の pptx を生成する: python3 test/fixtures/make_fixture.py

スライドの表示順と内部ファイル名（slideN.xml）がずれるよう、最初のスライドを末尾へ移動する。
"""
from pathlib import Path
from pptx import Presentation

prs = Presentation()
layout = prs.slide_layouts[5]

s1 = prs.slides.add_slide(layout)
s1.notes_slide.notes_text_frame.text = "\n".join([
    "冒頭のメモ",
    "```mermaid",
    "sequenceDiagram",
    "  User->>App: ログイン",
    "  App-->>User: OK",
    "```",
])

prs.slides.add_slide(layout)  # ノート無し

s3 = prs.slides.add_slide(layout)
# \v は段落内改行（Shift+Enter）として書き込まれる
s3.notes_slide.notes_text_frame.text = "```mermaid id=flow\vflowchart LR\v  A[“開始”] --> B\v```"

ids = prs.slides._sldIdLst
first = ids[0]
ids.remove(first)
ids.append(first)

out = Path(__file__).with_name("notes.pptx")
prs.save(out)
print(out)
